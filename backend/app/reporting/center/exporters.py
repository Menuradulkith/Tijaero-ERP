"""Report Center exports: CSV, Excel and PDF from one `ReportResult`.

A saved report's view (hidden columns, sort order per table) is applied first, so an
export matches what the user sees on screen.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Optional

from jinja2 import Environment, select_autoescape

from app.core import timezone as tz

from .core import Column, ReportDef, ReportResult, finalize_totals

_NUMERIC = ("int", "number", "money", "percent")


def apply_view(result: ReportResult, view: Optional[dict]) -> ReportResult:
    """Hide columns and sort rows as the user's view says (unknown keys are ignored)."""
    tables_view = (view or {}).get("tables") or {}
    for t in result.tables:
        tv = tables_view.get(t.key) or {}
        hidden = set(tv.get("hidden") or [])
        if hidden and len(hidden) < len(t.columns):
            t.columns = [c for c in t.columns if c.key not in hidden]
        sort = tv.get("sort") or {}
        col = next((c for c in t.columns if c.key == sort.get("key")), None)
        if col is not None:
            reverse = sort.get("dir") == "desc"
            present = [r for r in t.rows if r.get(col.key) is not None]
            missing = [r for r in t.rows if r.get(col.key) is None]
            key = (lambda r: Decimal(str(r[col.key]))) if col.type in _NUMERIC else (lambda r: str(r[col.key]).lower())
            t.rows = sorted(present, key=key, reverse=reverse) + missing
        finalize_totals(t)
    return result


def _display(value: Any, col: Column) -> str:
    if value is None:
        return ""
    if col.type == "bool":
        return "Yes" if value else "No"
    if col.type in ("money", "number"):
        return f"{Decimal(str(value)):,.2f}"
    if col.type == "percent":
        return f"{Decimal(str(value)):,.2f}%"
    if col.type == "int":
        return f"{int(value):,}"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _csv_safe(value: Any) -> Any:
    """Stop spreadsheet apps from treating text cells as formulas."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def to_csv(result: ReportResult, table_key: Optional[str] = None) -> bytes:
    t = next((t for t in result.tables if t.key == table_key), None) or (result.tables[0] if result.tables else None)
    buf = io.StringIO()
    writer = csv.writer(buf)
    if t is not None:
        writer.writerow([c.label for c in t.columns])
        for r in t.rows:
            writer.writerow([_csv_safe(_plain_cell(r.get(c.key), c)) for c in t.columns])
        if t.totals:
            writer.writerow([("Total" if i == 0 else _plain_cell(t.totals.get(c.key), c)) for i, c in enumerate(t.columns)])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel reads UTF-8


def _plain_cell(value: Any, col: Column) -> Any:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return float(value)
    if col.type == "bool":
        return "Yes" if value else "No"
    return value


def to_xlsx(report: ReportDef, params: dict, result: ReportResult, title: Optional[str] = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append([title or report.title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generated {tz.now():%Y-%m-%d %H:%M}"])
    ws.append([])
    for p in report.params:
        if params.get(p.key) is not None:
            ws.append([p.label, _xl(params[p.key])])
    if result.metrics:
        ws.append([])
        for m in result.metrics:
            ws.append([m.label, _xl(m.value)])
    for note in result.notes:
        ws.append([note])
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 22

    header_fill = PatternFill("solid", fgColor="1A4480")
    used: set[str] = {"Summary"}
    for t in result.tables:
        name = _sheet_name(t.title, used)
        sh = wb.create_sheet(name)
        sh.append([c.label for c in t.columns])
        for cell in sh[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for r in t.rows:
            sh.append([_xl(_csv_safe(r.get(c.key)) if c.type == "text" else r.get(c.key)) for c in t.columns])
        if t.totals:
            sh.append([("Total" if i == 0 else _xl(t.totals.get(c.key))) for i, c in enumerate(t.columns)])
            for cell in sh[sh.max_row]:
                cell.font = Font(bold=True)
        for i, c in enumerate(t.columns, start=1):
            letter = get_column_letter(i)
            fmt = {"money": "#,##0.00", "number": "#,##0.00", "int": "#,##0", "percent": '0.00"%"', "date": "yyyy-mm-dd"}.get(c.type)
            if fmt:
                for cell in sh[letter][1:]:
                    cell.number_format = fmt
            width = max([len(c.label)] + [len(_display(r.get(c.key), c)) for r in t.rows[:500]])
            sh.column_dimensions[letter].width = min(max(width + 2, 10), 60)
        sh.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _xl(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, str) and len(value) == 10 and value[4:5] == "-" and value[7:8] == "-":
        try:
            return date.fromisoformat(value)
        except ValueError:
            return value
    return value


def _sheet_name(title: str, used: set[str]) -> str:
    base = "".join(ch for ch in title if ch not in '[]:*?/\\')[:28] or "Data"
    name, n = base, 2
    while name in used:
        name = f"{base[:25]} {n}"
        n += 1
    used.add(name)
    return name


_env = Environment(autoescape=select_autoescape(default=True))
_HTML = _env.from_string("""<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>{{ title }}</title>
<style>
@page{size:A4 {{ 'landscape' if landscape else 'portrait' }};margin:12mm}
*{box-sizing:border-box}body{font-family:'Segoe UI',Arial,sans-serif;font-size:9px;color:#212529;margin:0}
.lh{display:flex;justify-content:space-between;align-items:flex-start;gap:16px;padding-bottom:8px;margin-bottom:10px;border-bottom:2px solid #1a4480}
.lh-co{display:flex;gap:10px;align-items:flex-start}.lh-logo{max-height:48px;max-width:130px;object-fit:contain}
.lh-name{font-size:13px;font-weight:700;color:#1a4480}.lh-sub{color:#6c757d;font-size:8px}
.lh-r{text-align:right}h1{font-size:16px;margin:0 0 2px;color:#1a4480}.sub{color:#6c757d;margin-bottom:8px}
.params{margin-bottom:8px}.params span{display:inline-block;margin-right:14px}.params b{color:#495057}
.metrics{display:flex;flex-wrap:wrap;gap:6px;margin-bottom:10px}
.metric{border:1px solid #dee2e6;border-radius:3px;padding:4px 8px;min-width:110px}
.metric .l{color:#6c757d;font-size:8px;text-transform:uppercase}.metric .v{font-size:12px;font-weight:600}
h2{font-size:11px;margin:12px 0 4px;color:#1a4480}
table{width:100%;border-collapse:collapse;page-break-inside:auto}tr{page-break-inside:avoid}
th{background:#1a4480;color:#fff;text-align:left;padding:3px 5px;font-weight:600}
td{border-bottom:1px solid #dee2e6;padding:3px 5px}td.n,th.n{text-align:right}
tfoot td{font-weight:700;border-top:2px solid #1a4480}.note{color:#6c757d;margin-top:6px}
</style></head><body>
<div class="lh">
  <div class="lh-co">
    {% if company.logo %}<img class="lh-logo" src="{{ company.logo }}" alt="">{% endif %}
    <div><div class="lh-name">{{ company.name }}</div>
      <div class="lh-sub">{{ company.address_lines|join(', ') }}</div>
      {% if company.tax_registration_number %}<div class="lh-sub">TIN / VAT No: {{ company.tax_registration_number }}</div>{% endif %}
    </div>
  </div>
  <div class="lh-r"><h1>{{ title }}</h1><div class="sub">Generated {{ generated }}</div></div>
</div>
{% if params %}<div class="params">{% for label, value in params %}<span><b>{{ label }}:</b> {{ value }}</span>{% endfor %}</div>{% endif %}
{% if metrics %}<div class="metrics">{% for m in metrics %}<div class="metric"><div class="l">{{ m.label }}</div><div class="v">{{ m.value }}</div></div>{% endfor %}</div>{% endif %}
{% for t in tables %}
<h2>{{ t.title }}</h2>
{% if t.rows %}<table><thead><tr>{% for c in t.columns %}<th class="{{ 'n' if c.numeric }}">{{ c.label }}</th>{% endfor %}</tr></thead>
<tbody>{% for r in t.rows %}<tr>{% for v in r %}<td class="{{ 'n' if t.columns[loop.index0].numeric }}">{{ v }}</td>{% endfor %}</tr>{% endfor %}</tbody>
{% if t.totals %}<tfoot><tr>{% for v in t.totals %}<td class="{{ 'n' if t.columns[loop.index0].numeric }}">{{ v }}</td>{% endfor %}</tr></tfoot>{% endif %}
</table>{% else %}<div class="note">No rows.</div>{% endif %}
{% endfor %}
{% for n in notes %}<div class="note">{{ n }}</div>{% endfor %}
</body></html>""")


def to_html(report: ReportDef, params: dict, result: ReportResult, company: dict, title: Optional[str] = None) -> tuple[str, bool]:
    """`company` is the letterhead profile from app.reporting.documents.company_profile."""
    """Printable HTML of a report, and whether it needs landscape paper."""
    widest = max((len(t.columns) for t in result.tables), default=0)
    tables = []
    for t in result.tables:
        cols = [{"label": c.label, "numeric": c.type in _NUMERIC} for c in t.columns]
        rows = [[_display(r.get(c.key), c) for c in t.columns] for r in t.rows]
        totals = (
            [("Total" if i == 0 else _display(t.totals.get(c.key), c)) for i, c in enumerate(t.columns)]
            if t.totals else None
        )
        tables.append({"title": t.title, "columns": cols, "rows": rows, "totals": totals})
    metric_cols = {"money": "money", "number": "number", "int": "int", "percent": "percent"}
    metrics = [
        {"label": m.label, "value": _display(m.value, Column("", "", metric_cols.get(m.type, "text")))}
        for m in result.metrics
    ]
    shown = [(p.label, _display(params.get(p.key), Column("", "", "text"))) for p in report.params if params.get(p.key) is not None]
    landscape = widest > 7
    html = _HTML.render(
        title=title or report.title,
        company=company,
        generated=f"{tz.now():%Y-%m-%d %H:%M}",
        params=shown,
        metrics=metrics,
        tables=tables,
        notes=result.notes,
        landscape=landscape,
    )
    return html, landscape


def to_pdf(html: str, landscape: bool) -> bytes:
    """Print the report HTML with headless Chromium (no JavaScript, no network)."""
    from app.reporting.chromium import chromium

    with chromium() as browser:
        page = browser.new_context(java_script_enabled=False).new_page()
        page.route("**/*", lambda route: route.continue_() if route.request.url.startswith("data:") else route.abort())
        page.set_content(html, wait_until="load")
        return page.pdf(
            format="A4",
            landscape=landscape,
            print_background=True,
            display_header_footer=True,
            header_template="<span></span>",
            footer_template='<div style="font-size:7px;width:100%;text-align:center;color:#999">'
                            'Page <span class="pageNumber"></span> of <span class="totalPages"></span></div>',
            margin={"top": "12mm", "bottom": "14mm", "left": "10mm", "right": "10mm"},
        )
